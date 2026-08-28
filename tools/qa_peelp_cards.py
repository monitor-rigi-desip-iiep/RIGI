#!/usr/bin/env python3
"""Regression checks for the PEELP modules and expanded project-card details."""

from __future__ import annotations

import json
import math
import subprocess
import sys
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

from qa_data_contract import validate_workbook


ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "data/RIGI_tracker_data_final_con_proyectos_integrados.xlsx"


def read(relative: str) -> str:
    return (ROOT / relative).read_text(encoding="utf-8")


def normalized(value: object) -> str:
    text = "" if value is None else str(value)
    text = " ".join(text.strip().split()).casefold()
    return "".join(
        char for char in unicodedata.normalize("NFD", text)
        if unicodedata.category(char) != "Mn"
    )


def as_number(value: object) -> float | None:
    if value is None or value == "":
        return None
    try:
        parsed = float(value)
    except (TypeError, ValueError):
        return None
    return parsed if math.isfinite(parsed) else None


data_contract = validate_workbook(XLSX)
if not data_contract["valid"]:
    report = {
        "source_sha256": data_contract["sha256"],
        "checks": {"source_xlsx_contract_valid": False},
        "data_contract": data_contract,
        "failed": ["source_xlsx_contract_valid"],
    }
    (ROOT / "qa").mkdir(exist_ok=True)
    serialized = json.dumps(report, ensure_ascii=False, indent=2) + "\n"
    (ROOT / "qa/peelp_cards_source_qa.json").write_text(serialized, encoding="utf-8")
    print(serialized, end="")
    raise SystemExit(1)

workbook = load_workbook(XLSX, read_only=True, data_only=True)
projects_sheet = workbook["Proyectos"]
rows = list(projects_sheet.iter_rows(values_only=True))
headers = [str(value) if value is not None else "" for value in rows[0]]
records = [
    dict(zip(headers, row))
    for row in rows[1:]
    if any(value is not None for value in row)
]

approved = [
    row for row in records
    if normalized(row.get("estado_administrativo")) == "aprobado"
]
peelp = [row for row in approved if normalized(row.get("peelp")) in {"si", "true", "1"}]

approved_amounts = [
    value for row in approved
    if (value := as_number(row.get("inversion_total_mill_usd"))) is not None
]
peelp_amounts = [
    value for row in peelp
    if (value := as_number(row.get("inversion_total_mill_usd"))) is not None
]
approved_total = sum(approved_amounts)
peelp_total = sum(peelp_amounts)
non_peelp_total = approved_total - peelp_total
share_sum = (
    peelp_total / approved_total + non_peelp_total / approved_total
    if approved_total > 0 else 0.0
)

approved_qmd = read("aprobados.qmd")
plots = read("R/04_plots.R")
modules = read("R/07_investment_modules.R")
javascript = read("assets/investment_modules.js")
styles = read("styles.css")
workflow = read(".github/workflows/render.yml")

card_start = plots.index("make_rigi_project_card <- function")
card_end = plots.index("make_rigi_project_cards <- function")
card_source = plots[card_start:card_end]
details_start = card_source.index("detail_fields <-")
details_end = card_source.index("htmltools::tags$article(", details_start)
details_source = card_source[details_start:details_end]
header_start = card_source.index('class = "rigi-project-card__header"')
header_end = card_source.index('class = "rigi-project-card__summary-grid"', header_start)
header_source = card_source[header_start:header_end]

detail_labels = [
    "Inversión total",
    "Inversión en activos computables",
    "Inversión comprometida en activos computables — primeros 2 años",
]
detail_positions = [details_source.find(f'"{label}"') for label in detail_labels]

checks = {
    "source_xlsx_contract_valid": data_contract["valid"],
    "peelp_universe_is_approved_only": all(row in approved for row in peelp),
    "peelp_and_non_peelp_reconcile": math.isclose(
        peelp_total + non_peelp_total, approved_total, rel_tol=0, abs_tol=1e-9
    ),
    "peelp_shares_sum_to_one": approved_total <= 0
    or math.isclose(share_sum, 1.0, rel_tol=0, abs_tol=1e-12),
    "peelp_ranking_descends_by_amount": peelp_amounts == sorted(peelp_amounts, reverse=True)
    or "dplyr::arrange(dplyr::desc(value), label)" in modules,
    "peelp_share_module_used": "make_peelp_share_module(indicadores)" in approved_qmd,
    "peelp_ranking_module_used": "make_peelp_ranking_module(" in approved_qmd,
    "old_peelp_pie_removed": all(
        token not in plots + "\n" + approved_qmd
        for token in ("plot_peelp_share <-", "plot_peelp_share(indicadores)", 'type = "pie"')
    ),
    "peelp_share_uses_approved_denominator": all(
        token in modules
        for token in (
            "approved_amount - peelp_amount",
            "ratio_or_na(rows$value[[index]], approved_amount)",
            "100 * share",
        )
    ),
    "peelp_ranking_uses_approved_denominator": (
        "share_total = approved_total" in modules
        and "approved_total = indicadores$monto_aprobado" in approved_qmd
    ),
    "peelp_ranking_defaults_to_five": "default_limit = 5L" in modules,
    "peelp_selector_options": all(
        token in modules
        for token in ('value = "5"', 'value = "10"', 'value = "15"', 'value = "all"', '"Todos"')
    ),
    "peelp_ranking_has_no_empty_tabs": "show_tabs = FALSE" in modules,
    "peelp_modules_accessible": all(
        token in modules
        for token in (
            '`data-peelp-share-module` = "true"',
            'tabindex = "0"',
            '`aria-label` = accessible',
            '`aria-live` = "polite"',
        )
    ),
    "peelp_responsive_styles_present": all(
        token in styles for token in (".rigi-peelp-share-module", "#rigi-peelp-ranking", "@media (max-width: 640px)")
    ),
    "single_header_currency_metric": header_source.count('htmltools::tags$span("Inversión total")') == 1,
    "extra_metrics_absent_from_header": all(
        label not in header_source for label in detail_labels[1:]
    ),
    "three_detail_metrics_present_once": all(
        details_source.count(f'"{label}"') == 1 for label in detail_labels
    ),
    "detail_metrics_in_requested_order": all(position >= 0 for position in detail_positions)
    and detail_positions == sorted(detail_positions),
    "details_use_no_informed_for_missing": 'return("No informado")' in plots
    and "rigi_card_amount" in details_source,
    "zero_is_not_treated_as_missing": "is.na(x) || is.nan(x)" in plots
    and "x == 0" not in read("R/04_plots.R")[read("R/04_plots.R").index("rigi_card_amount <- function"):card_start],
    "shared_ranked_javascript_reused": all(
        token in javascript for token in ("initRankedModule", "config.shareTotal", "config.defaultLimit")
    ),
    "workflow_runs_peelp_qa": "python3 tools/qa_peelp_cards.py" in workflow,
}

javascript_results: dict[str, dict[str, object]] = {}
for path in sorted((ROOT / "assets").glob("*.js")):
    result = subprocess.run(
        ["node", "--check", str(path)],
        capture_output=True,
        text=True,
        check=False,
    )
    key = str(path.relative_to(ROOT))
    javascript_results[key] = {
        "ok": result.returncode == 0,
        "stderr": result.stderr.strip(),
    }
    checks[f"javascript_syntax_{path.stem}"] = result.returncode == 0

failed = sorted(name for name, passed in checks.items() if not passed)
report = {
    "source_sha256": data_contract["sha256"],
    "data_contract": {
        "metrics": data_contract["metrics"],
        "errors": data_contract["errors"],
    },
    "projects": len(records),
    "approved_projects": len(approved),
    "peelp_projects": len(peelp),
    "approved_total_mill_usd": approved_total,
    "peelp_total_mill_usd": peelp_total,
    "non_peelp_total_mill_usd": non_peelp_total,
    "share_sum": share_sum,
    "checks": checks,
    "javascript": javascript_results,
    "failed": failed,
}
(ROOT / "qa").mkdir(exist_ok=True)
serialized = json.dumps(report, ensure_ascii=False, indent=2) + "\n"
(ROOT / "qa/peelp_cards_source_qa.json").write_text(serialized, encoding="utf-8")
print(serialized, end="")
sys.exit(1 if failed else 0)
