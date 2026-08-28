#!/usr/bin/env python3
"""Pruebas de que el contrato acepta datos nuevos y rechaza esquemas rotos."""

from __future__ import annotations

import json
import shutil
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from qa_data_contract import DEFAULT_XLSX, ROOT, validate_workbook


def header_column(worksheet: object, header: str) -> int:
    for cell in worksheet[1]:
        if str(cell.value).strip() == header:
            return int(cell.column)
    raise RuntimeError(f"No se encontró la columna de prueba: {header}")


def main() -> int:
    with tempfile.TemporaryDirectory(prefix="rigi-data-contract-") as directory:
        temporary = Path(directory)

        valid_update = temporary / "valid_update.xlsx"
        shutil.copy2(DEFAULT_XLSX, valid_update)
        workbook = load_workbook(valid_update)
        projects = workbook["Proyectos"]
        updated_column = header_column(projects, "vpu")
        for row in range(2, projects.max_row + 1):
            cell = projects.cell(row=row, column=updated_column)
            if isinstance(cell.value, str) and cell.value.strip():
                cell.value = cell.value + " [prueba temporal de actualización válida]"
                break
        workbook.save(valid_update)
        workbook.close()
        valid_result = validate_workbook(valid_update)

        broken_schema = temporary / "broken_schema.xlsx"
        shutil.copy2(DEFAULT_XLSX, broken_schema)
        workbook = load_workbook(broken_schema)
        projects = workbook["Proyectos"]
        projects.delete_cols(header_column(projects, "id_proyecto"), 1)
        workbook.save(broken_schema)
        workbook.close()
        broken_result = validate_workbook(broken_schema)

    checks = {
        "valid_content_update_is_accepted": valid_result["valid"],
        "changed_sha_is_accepted": valid_result["sha256"] != validate_workbook(DEFAULT_XLSX)["sha256"],
        "missing_required_column_is_rejected": not broken_result["valid"],
        "missing_column_error_is_actionable": any(
            error.get("check") == "source_xlsx_required_columns"
            and error.get("column") == "id_proyecto"
            for error in broken_result["errors"]
        ),
    }
    report = {
        "checks": checks,
        "valid_update": {
            "valid": valid_result["valid"],
            "sha256": valid_result["sha256"],
            "failed": valid_result["failed"],
        },
        "broken_schema": {
            "valid": broken_result["valid"],
            "failed": broken_result["failed"],
            "errors": broken_result["errors"],
        },
        "failed": sorted(name for name, passed in checks.items() if not passed),
    }
    output = ROOT / "qa/data_contract_self_test.json"
    output.parent.mkdir(exist_ok=True)
    serialized = json.dumps(report, ensure_ascii=False, indent=2) + "\n"
    output.write_text(serialized, encoding="utf-8")
    print(serialized, end="")
    return 1 if report["failed"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
