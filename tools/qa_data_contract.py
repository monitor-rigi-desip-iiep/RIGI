#!/usr/bin/env python3
"""Valida el contrato estable del Excel maestro sin fijar sus datos ni su hash."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import sys
import unicodedata
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = ROOT / "data/RIGI_tracker_data_final_con_proyectos_integrados.xlsx"
DEFAULT_OUTPUT = ROOT / "qa/data_contract_qa.json"

REQUIRED_SHEETS = ("Proyectos", "Diccionario", "Variable_Anterior")
REQUIRED_PROJECT_COLUMNS = (
    "vpu",
    "descripcion_proyecto",
    "peelp",
    "id_proyecto",
    "empresa",
    "titular_proyecto",
    "cuit",
    "sector",
    "subsector",
    "actividad_subsector_resolucion_mecon",
    "provincia",
    "localidad_region",
    "inversion_total_mill_usd",
    "inversion_total_resolucion_mill_usd",
    "inversion_activos_computables_mill_usd",
    "inversion_activos_computables_comprometida_anio_1_mill_usd",
    "inversion_activos_computables_comprometida_anio_2_mill_usd",
    "inversion_activos_computables_comprometida_2_anios_mill_usd",
    "estado_administrativo",
    "empleos_directos_indirectos_informados",
    "fecha_limite_inversion_minima_activos_computables",
    "fecha_presentacion",
    "fecha_adhesion_rigi",
    "fecha_publicacion_bo",
    "norma_aprobacion",
    "clasificacion_preexistencia_boletin_oficial",
    "justificacion_preexistencia_boletin_oficial",
    "link_norma",
    "fuentes",
    "links_fuentes",
)
REQUIRED_DICTIONARY_COLUMNS = (
    "orden",
    "variable",
    "nombre_visible",
    "tipo",
    "unidad",
    "descripcion_breve",
)
REQUIRED_MAPPING_COLUMNS = (
    "orden",
    "nombre_actual_en_xlsx_madre",
    "variable_recomendada",
)
NUMERIC_COLUMNS = (
    "inversion_total_mill_usd",
    "inversion_total_resolucion_mill_usd",
    "inversion_activos_computables_mill_usd",
    "inversion_activos_computables_comprometida_anio_1_mill_usd",
    "inversion_activos_computables_comprometida_anio_2_mill_usd",
    "inversion_activos_computables_comprometida_2_anios_mill_usd",
    "empleos_directos_indirectos_informados",
)
DATE_COLUMNS = (
    "fecha_limite_inversion_minima_activos_computables",
    "fecha_presentacion",
    "fecha_adhesion_rigi",
    "fecha_publicacion_bo",
)
ALLOWED_STATES = {
    "aprobado": "Aprobado",
    "en evaluacion": "En evaluación",
    "rechazado": "Rechazado",
}
CORE_REQUIRED_COLUMNS = ("vpu", "estado_administrativo")
MISSING_MARKERS = {
    "-",
    "—",
    "n/a",
    "n/d",
    "na",
    "nd",
    "no informado",
    "s/d",
    "sin informar",
}
APPROVED_REQUIRED_COLUMNS = (
    "inversion_total_mill_usd",
    "fecha_adhesion_rigi",
    "fecha_publicacion_bo",
    "norma_aprobacion",
    "link_norma",
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_text(value: object) -> str:
    text = "" if value is None else " ".join(str(value).strip().split()).casefold()
    decomposed = unicodedata.normalize("NFD", text)
    return "".join(char for char in decomposed if unicodedata.category(char) != "Mn")


def normalize_header(value: object) -> str:
    normalized = normalize_text(value)
    normalized = re.sub(r"[^a-z0-9]+", "_", normalized)
    return normalized.strip("_")


def is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def is_missing_value(value: object) -> bool:
    return is_blank(value) or (
        isinstance(value, str) and normalize_text(value) in MISSING_MARKERS
    )


def canonical_id(value: object) -> str | None:
    if is_missing_value(value):
        return None
    if isinstance(value, bool):
        return str(value).casefold()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip().casefold()


def is_numeric_compatible(value: object) -> bool:
    if is_missing_value(value):
        return True
    if isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        return math.isfinite(float(value))
    if not isinstance(value, str):
        return False
    candidate = value.strip().replace("\u00a0", "").replace(" ", "")
    if not candidate:
        return True
    if "," in candidate and "." in candidate:
        if candidate.rfind(",") > candidate.rfind("."):
            candidate = candidate.replace(".", "").replace(",", ".")
        else:
            candidate = candidate.replace(",", "")
    elif "," in candidate:
        candidate = candidate.replace(",", ".")
    try:
        return math.isfinite(float(candidate))
    except ValueError:
        return False


def is_date_compatible(value: object) -> bool:
    if is_missing_value(value):
        return True
    if isinstance(value, bool):
        return False
    if isinstance(value, (datetime, date)):
        return True
    if isinstance(value, (int, float)):
        return math.isfinite(float(value)) and float(value) > 0
    if not isinstance(value, str):
        return False
    candidate = value.strip()
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            datetime.strptime(candidate, pattern)
            return True
        except ValueError:
            continue
    try:
        datetime.fromisoformat(candidate.replace("Z", "+00:00"))
        return True
    except ValueError:
        return False


def utf8_compatible(value: object) -> bool:
    if not isinstance(value, str):
        return True
    if "\ufffd" in value:
        return False
    try:
        value.encode("utf-8", errors="strict")
        return True
    except UnicodeError:
        return False


def workbook_records(worksheet: Any) -> tuple[list[str], list[dict[str, object]], list[int]]:
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return [], [], []
    headers = [normalize_header(value) for value in rows[0]]
    nonempty_positions = [
        index
        for index, row in enumerate(rows[1:], start=2)
        if any(not is_blank(value) for value in row)
    ]
    last_nonempty = max(nonempty_positions, default=1)
    blank_rows = [
        index
        for index, row in enumerate(rows[1:last_nonempty], start=2)
        if all(is_blank(value) for value in row)
    ]
    records = [
        dict(zip(headers, row))
        for index, row in enumerate(rows[1:], start=2)
        if index <= last_nonempty and any(not is_blank(value) for value in row)
    ]
    return headers, records, blank_rows


def validate_workbook(path: Path | str = DEFAULT_XLSX) -> dict[str, Any]:
    path = Path(path).resolve()
    check_names = (
        "source_xlsx_exists",
        "source_xlsx_nonempty",
        "source_xlsx_container_valid",
        "source_xlsx_readable",
        "source_xlsx_required_sheets",
        "source_xlsx_project_headers_unique",
        "source_xlsx_required_columns",
        "source_xlsx_no_unexpected_columns",
        "source_xlsx_has_projects",
        "source_xlsx_no_internal_blank_rows",
        "source_xlsx_project_ids_present",
        "source_xlsx_project_ids_unique",
        "source_xlsx_core_fields_present",
        "source_xlsx_valid_statuses",
        "source_xlsx_numeric_compatibility",
        "source_xlsx_date_compatibility",
        "source_xlsx_utf8_compatibility",
        "source_xlsx_approved_fields_complete",
        "source_xlsx_dictionary_contract",
        "source_xlsx_mapping_contract",
    )
    checks: dict[str, bool | None] = {name: None for name in check_names}
    details: dict[str, Any] = {
        "path": str(path),
        "sha256": None,
        "errors": [],
        "metrics": {},
    }

    def error(
        check: str,
        message: str,
        *,
        sheet: str | None = None,
        column: str | None = None,
        row: int | None = None,
        value: object | None = None,
    ) -> None:
        item: dict[str, Any] = {"check": check, "message": message}
        if sheet is not None:
            item["sheet"] = sheet
        if column is not None:
            item["column"] = column
        if row is not None:
            item["row"] = row
        if value is not None:
            item["value"] = str(value)
        details["errors"].append(item)

    checks["source_xlsx_exists"] = path.is_file()
    if not checks["source_xlsx_exists"]:
        error("source_xlsx_exists", f"No se encontró el Excel maestro: {path}")
        return finalize(checks, details)

    checks["source_xlsx_nonempty"] = path.stat().st_size > 0
    if not checks["source_xlsx_nonempty"]:
        error("source_xlsx_nonempty", f"El Excel maestro está vacío: {path}")
        return finalize(checks, details)

    details["sha256"] = sha256(path)
    checks["source_xlsx_container_valid"] = zipfile.is_zipfile(path)
    if not checks["source_xlsx_container_valid"]:
        error(
            "source_xlsx_container_valid",
            "El archivo no es un contenedor XLSX válido o está truncado.",
        )
        return finalize(checks, details)

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl expone varias excepciones según el daño
        error("source_xlsx_readable", f"No se pudo abrir el XLSX: {exc}")
        return finalize(checks, details)

    checks["source_xlsx_readable"] = True
    missing_sheets = [name for name in REQUIRED_SHEETS if name not in workbook.sheetnames]
    checks["source_xlsx_required_sheets"] = not missing_sheets
    for sheet in missing_sheets:
        error(
            "source_xlsx_required_sheets",
            f"Falta la hoja obligatoria: {sheet}",
            sheet=sheet,
        )
    if missing_sheets:
        workbook.close()
        return finalize(checks, details)

    project_headers, records, blank_rows = workbook_records(workbook["Proyectos"])
    duplicated_headers = sorted(
        {header for header in project_headers if header and project_headers.count(header) > 1}
    )
    blank_headers = [index + 1 for index, header in enumerate(project_headers) if not header]
    checks["source_xlsx_project_headers_unique"] = not duplicated_headers and not blank_headers
    if duplicated_headers:
        error(
            "source_xlsx_project_headers_unique",
            "La hoja Proyectos contiene encabezados duplicados: "
            + ", ".join(duplicated_headers),
            sheet="Proyectos",
        )
    if blank_headers:
        error(
            "source_xlsx_project_headers_unique",
            "La hoja Proyectos contiene encabezados vacíos en las columnas: "
            + ", ".join(map(str, blank_headers)),
            sheet="Proyectos",
        )

    missing_columns = [name for name in REQUIRED_PROJECT_COLUMNS if name not in project_headers]
    unexpected_columns = [
        name for name in project_headers if name and name not in REQUIRED_PROJECT_COLUMNS
    ]
    checks["source_xlsx_required_columns"] = not missing_columns
    checks["source_xlsx_no_unexpected_columns"] = not unexpected_columns
    for column in missing_columns:
        error(
            "source_xlsx_required_columns",
            f"La hoja Proyectos no contiene la columna obligatoria: {column}",
            sheet="Proyectos",
            column=column,
        )
    if unexpected_columns:
        error(
            "source_xlsx_no_unexpected_columns",
            "La hoja Proyectos contiene columnas no previstas por el pipeline de R: "
            + ", ".join(unexpected_columns),
            sheet="Proyectos",
        )

    checks["source_xlsx_has_projects"] = bool(records)
    if not records:
        error(
            "source_xlsx_has_projects",
            "La hoja Proyectos no contiene proyectos.",
            sheet="Proyectos",
        )

    checks["source_xlsx_no_internal_blank_rows"] = not blank_rows
    if blank_rows:
        error(
            "source_xlsx_no_internal_blank_rows",
            "La tabla contiene filas completamente vacías antes del último proyecto: "
            + ", ".join(map(str, blank_rows[:20])),
            sheet="Proyectos",
        )

    if missing_columns or duplicated_headers or blank_headers:
        workbook.close()
        return finalize(checks, details)

    missing_ids: list[int] = []
    seen_ids: dict[str, int] = {}
    duplicated_ids: list[tuple[str, int, int]] = []
    missing_core_fields: list[tuple[int, str, object]] = []
    invalid_statuses: list[tuple[int, object]] = []
    invalid_numeric: list[tuple[int, str, object]] = []
    invalid_dates: list[tuple[int, str, object]] = []
    invalid_utf8: list[tuple[str, int, str, object]] = []
    incomplete_approved: list[tuple[int, str, object]] = []
    state_counts = {label: 0 for label in ALLOWED_STATES.values()}

    for offset, record in enumerate(records, start=2):
        project_id = canonical_id(record.get("id_proyecto"))
        if project_id is None:
            missing_ids.append(offset)
        elif project_id in seen_ids:
            duplicated_ids.append((project_id, seen_ids[project_id], offset))
        else:
            seen_ids[project_id] = offset

        for column in CORE_REQUIRED_COLUMNS:
            if is_missing_value(record.get(column)):
                missing_core_fields.append((offset, column, record.get("id_proyecto")))

        state_key = normalize_text(record.get("estado_administrativo"))
        if state_key not in ALLOWED_STATES:
            invalid_statuses.append((offset, record.get("estado_administrativo")))
        else:
            state_counts[ALLOWED_STATES[state_key]] += 1

        for column in NUMERIC_COLUMNS:
            value = record.get(column)
            if not is_numeric_compatible(value):
                invalid_numeric.append((offset, column, value))

        for column in DATE_COLUMNS:
            value = record.get(column)
            if not is_date_compatible(value):
                invalid_dates.append((offset, column, value))

        for column, value in record.items():
            if not utf8_compatible(value):
                invalid_utf8.append(("Proyectos", offset, column, value))

        if state_key == "aprobado":
            for column in APPROVED_REQUIRED_COLUMNS:
                if is_missing_value(record.get(column)):
                    incomplete_approved.append((offset, column, record.get("id_proyecto")))

    checks["source_xlsx_project_ids_present"] = not missing_ids
    checks["source_xlsx_project_ids_unique"] = not duplicated_ids
    checks["source_xlsx_core_fields_present"] = not missing_core_fields
    checks["source_xlsx_valid_statuses"] = not invalid_statuses
    checks["source_xlsx_numeric_compatibility"] = not invalid_numeric
    checks["source_xlsx_date_compatibility"] = not invalid_dates
    checks["source_xlsx_approved_fields_complete"] = not incomplete_approved

    for row in missing_ids[:20]:
        error(
            "source_xlsx_project_ids_present",
            "La fila de proyecto no tiene id_proyecto.",
            sheet="Proyectos",
            column="id_proyecto",
            row=row,
        )
    for project_id, first_row, row in duplicated_ids[:20]:
        error(
            "source_xlsx_project_ids_unique",
            f"id_proyecto duplicado; también aparece en la fila {first_row}.",
            sheet="Proyectos",
            column="id_proyecto",
            row=row,
            value=project_id,
        )
    for row, column, project_id in missing_core_fields[:20]:
        error(
            "source_xlsx_core_fields_present",
            "El proyecto no contiene un campo mínimo obligatorio.",
            sheet="Proyectos",
            column=column,
            row=row,
            value=f"id_proyecto={project_id}",
        )
    for row, value in invalid_statuses[:20]:
        error(
            "source_xlsx_valid_statuses",
            "Estado administrativo no admitido. Valores permitidos: "
            + ", ".join(ALLOWED_STATES.values()),
            sheet="Proyectos",
            column="estado_administrativo",
            row=row,
            value=value,
        )
    for row, column, value in invalid_numeric[:20]:
        error(
            "source_xlsx_numeric_compatibility",
            "El valor no puede interpretarse como número ni como dato faltante.",
            sheet="Proyectos",
            column=column,
            row=row,
            value=value,
        )
    for row, column, value in invalid_dates[:20]:
        error(
            "source_xlsx_date_compatibility",
            "El valor no puede interpretarse como fecha ni como dato faltante.",
            sheet="Proyectos",
            column=column,
            row=row,
            value=value,
        )
    for row, column, project_id in incomplete_approved[:20]:
        error(
            "source_xlsx_approved_fields_complete",
            "Un proyecto aprobado no contiene un campo obligatorio para su publicación.",
            sheet="Proyectos",
            column=column,
            row=row,
            value=f"id_proyecto={project_id}",
        )

    dictionary_headers, dictionary_rows, _ = workbook_records(workbook["Diccionario"])
    missing_dictionary_headers = [
        name for name in REQUIRED_DICTIONARY_COLUMNS if name not in dictionary_headers
    ]
    dictionary_variables = [
        normalize_header(row.get("variable"))
        for row in dictionary_rows
        if not is_blank(row.get("variable"))
    ]
    missing_definitions = [
        name for name in REQUIRED_PROJECT_COLUMNS if name not in dictionary_variables
    ]
    duplicated_definitions = sorted(
        {name for name in dictionary_variables if dictionary_variables.count(name) > 1}
    )
    checks["source_xlsx_dictionary_contract"] = not (
        missing_dictionary_headers or missing_definitions or duplicated_definitions
    )
    if missing_dictionary_headers:
        error(
            "source_xlsx_dictionary_contract",
            "La hoja Diccionario no contiene las columnas obligatorias: "
            + ", ".join(missing_dictionary_headers),
            sheet="Diccionario",
        )
    if missing_definitions:
        error(
            "source_xlsx_dictionary_contract",
            "El Diccionario no define las variables: " + ", ".join(missing_definitions),
            sheet="Diccionario",
            column="Variable",
        )
    if duplicated_definitions:
        error(
            "source_xlsx_dictionary_contract",
            "El Diccionario contiene variables duplicadas: "
            + ", ".join(duplicated_definitions),
            sheet="Diccionario",
            column="Variable",
        )

    mapping_headers, mapping_rows, _ = workbook_records(workbook["Variable_Anterior"])
    missing_mapping_headers = [
        name for name in REQUIRED_MAPPING_COLUMNS if name not in mapping_headers
    ]
    mapping_targets = [
        normalize_header(row.get("variable_recomendada"))
        for row in mapping_rows
        if not is_blank(row.get("variable_recomendada"))
    ]
    missing_targets = [name for name in REQUIRED_PROJECT_COLUMNS if name not in mapping_targets]
    duplicated_targets = sorted(
        {name for name in mapping_targets if mapping_targets.count(name) > 1}
    )
    checks["source_xlsx_mapping_contract"] = not (
        missing_mapping_headers or missing_targets or duplicated_targets
    )
    if missing_mapping_headers:
        error(
            "source_xlsx_mapping_contract",
            "La hoja Variable_Anterior no contiene las columnas obligatorias: "
            + ", ".join(missing_mapping_headers),
            sheet="Variable_Anterior",
        )
    if missing_targets:
        error(
            "source_xlsx_mapping_contract",
            "Variable_Anterior no documenta las variables: " + ", ".join(missing_targets),
            sheet="Variable_Anterior",
            column="Variable recomendada",
        )
    if duplicated_targets:
        error(
            "source_xlsx_mapping_contract",
            "Variable_Anterior contiene destinos duplicados: "
            + ", ".join(duplicated_targets),
            sheet="Variable_Anterior",
            column="Variable recomendada",
        )

    for sheet_name in REQUIRED_SHEETS:
        headers, sheet_rows, _ = workbook_records(workbook[sheet_name])
        for row_number, record in enumerate(sheet_rows, start=2):
            for column in headers:
                value = record.get(column)
                if not utf8_compatible(value):
                    invalid_utf8.append((sheet_name, row_number, column, value))
    checks["source_xlsx_utf8_compatibility"] = not invalid_utf8
    for sheet, row, column, value in invalid_utf8[:20]:
        error(
            "source_xlsx_utf8_compatibility",
            "El texto no es UTF-8 válido o contiene el carácter de reemplazo.",
            sheet=sheet,
            column=column,
            row=row,
            value=value,
        )

    details["metrics"] = {
        "projects": len(records),
        "states": state_counts,
        "sheets": list(workbook.sheetnames),
        "project_columns": project_headers,
    }
    workbook.close()
    return finalize(checks, details)


def finalize(checks: dict[str, bool | None], details: dict[str, Any]) -> dict[str, Any]:
    failed = sorted(name for name, passed in checks.items() if passed is False)
    skipped = sorted(name for name, passed in checks.items() if passed is None)
    return {
        "valid": not failed,
        "checks": checks,
        "sha256": details.get("sha256"),
        "metrics": details.get("metrics", {}),
        "errors": details.get("errors", []),
        "failed": failed,
        "skipped": skipped,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", default=str(DEFAULT_XLSX))
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT))
    parser.add_argument("--no-write", action="store_true")
    args = parser.parse_args()

    result = validate_workbook(args.xlsx)
    serialized = json.dumps(result, ensure_ascii=False, indent=2) + "\n"
    if not args.no_write:
        output = Path(args.output)
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(serialized, encoding="utf-8")
    print(serialized, end="")
    return 0 if result["valid"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
